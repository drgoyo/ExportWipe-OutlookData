"""Generate genuine files and verify Office/table/text/image parsing outcomes."""

import shutil
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from pipelines.document_parsing.parser import MarkdownParser
from project_name.core.records import ParseOptions


def test_word_excel_powerpoint_csv_and_text(tmp_path: Path) -> None:
    word = Document()
    word.add_heading("Renewal contract", level=1)
    word.add_paragraph("The service renews in 2026.")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Service"
    table.cell(1, 1).text = "1200"
    word.save(str(tmp_path / "contract.docx"))
    workbook = Workbook()
    worksheet = workbook.create_sheet("Values")
    worksheet.append(["Name", "Amount"])
    worksheet.append(["Service", 1200])
    workbook.save(tmp_path / "values.xlsx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "Contract renewal"
    slide.placeholders[1].text = "Renewal date: 2026"
    slides.save(str(tmp_path / "slides.pptx"))
    (tmp_path / "values.csv").write_text('name,amount\n"Service | plan",1200\n')
    (tmp_path / "body.txt").write_text("Hello\r\n\r\n\r\nRenewal\x00", encoding="utf-8")
    parser = MarkdownParser()
    assert "# Renewal contract" in parser.parse(tmp_path / "contract.docx")
    assert "| Service | 1200 |" in parser.parse(tmp_path / "values.xlsx")
    assert "## Slide 1" in parser.parse(tmp_path / "slides.pptx")
    assert "Service \\| plan" in parser.parse(tmp_path / "values.csv")
    assert parser.parse(tmp_path / "body.txt") == "Hello\n\nRenewal\n"


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="External Tesseract is not installed")
def test_real_jpeg_ocr(tmp_path: Path) -> None:
    image = Image.new("RGB", (1000, 140), "white")
    ImageDraw.Draw(image).text(
        (25, 30), "Contract renewal 2026", font=ImageFont.load_default(size=48), fill="black"
    )
    path = tmp_path / "page.jpg"
    image.save(path, quality=95)
    assert "Contract renewal" in MarkdownParser().parse(path)


def test_size_format_and_table_bounds(tmp_path: Path) -> None:
    text = tmp_path / "large.txt"
    text.write_text("too much text")
    with pytest.raises(ValueError, match="size limit"):
        MarkdownParser(ParseOptions(max_file_bytes=2)).parse(text)
    unknown = tmp_path / "document.exe"
    unknown.write_text("not a document")
    with pytest.raises(ValueError, match="Unsupported"):
        MarkdownParser().parse(unknown)
    csv = tmp_path / "rows.csv"
    csv.write_text("value\n1\n2\n3\n")
    with pytest.raises(ValueError, match="row limit"):
        MarkdownParser(ParseOptions(max_table_rows=1)).parse(csv)


def test_corrupt_office_file_is_a_reportable_parse_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip archive")
    with pytest.raises(ValueError, match="Cannot parse"):
        MarkdownParser().parse(path)
