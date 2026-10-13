"""Bounded, deterministic Office/text/image conversion to clean Markdown."""

from __future__ import annotations

import re
from typing import TYPE_CHECKING
from zipfile import ZipFile

import polars as pl

from project_name.core.records import ParseOptions

if TYPE_CHECKING:
    from pathlib import Path

SUPPORTED_EXTENSIONS = {".docx", ".xlsx", ".pptx", ".csv", ".txt", ".md", ".jpg", ".jpeg", ".png"}


def clean_markdown(value: str) -> str:
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = "".join(character for character in value if ord(character) >= 32 or character in "\n\t")
    return (
        re.sub(r"\n{3,}", "\n\n", "\n".join(line.rstrip() for line in value.splitlines())).strip()
        + "\n"
    )


def table_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(map(len, rows))
    escaped = [
        [cell.replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>") for cell in row]
        + [""] * (width - len(row))
        for row in rows
    ]
    lines = ["| " + " | ".join(escaped[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in escaped[1:])
    return "\n".join(lines)


class MarkdownParser:
    def __init__(self, options: ParseOptions | None = None) -> None:
        self.options = options or ParseOptions()

    def parse(self, path: Path) -> str:
        try:
            return self._parse(path)
        except (ValueError, OSError, RuntimeError):
            raise
        except Exception as exc:
            # A corrupt third-party document must not abort unrelated manifest records.
            raise ValueError(f"Cannot parse {path.name}: {type(exc).__name__}") from None

    def _parse(self, path: Path) -> str:
        extension = path.suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported document format: {extension}")
        if not path.is_file() or path.stat().st_size > self.options.max_file_bytes:
            raise ValueError(f"File missing or exceeds parsing size limit: {path.name}")
        if extension in {".docx", ".xlsx", ".pptx"}:
            with ZipFile(path) as archive:
                if (
                    sum(item.file_size for item in archive.infolist())
                    > self.options.max_uncompressed_bytes
                ):
                    raise ValueError("Office archive exceeds expanded-size limit")
        handlers = {
            ".docx": self.word,
            ".xlsx": self.excel,
            ".pptx": self.powerpoint,
            ".csv": self.csv,
            ".txt": self.text,
            ".md": self.text,
            ".jpg": self.image,
            ".jpeg": self.image,
            ".png": self.image,
        }
        result = clean_markdown(handlers[extension](path))
        if not result.strip():
            raise ValueError(f"No readable content extracted from {path.name}")
        return result

    def text(self, path: Path) -> str:
        return path.read_text(encoding="utf-8-sig", errors="strict")

    def csv(self, path: Path) -> str:
        frame = pl.read_csv(path, infer_schema=False, n_rows=self.options.max_table_rows + 1)
        if frame.height > self.options.max_table_rows:
            raise ValueError("CSV exceeds table row limit")
        rows = [
            frame.columns,
            *[["" if value is None else str(value) for value in row] for row in frame.iter_rows()],
        ]
        return table_markdown(rows)

    def word(self, path: Path) -> str:
        from docx import Document as WordDocument
        from docx.table import Table

        result = []
        for block in WordDocument(str(path)).iter_inner_content():
            if isinstance(block, Table):
                if len(block.rows) > self.options.max_table_rows:
                    raise ValueError("Word table exceeds row limit")
                result.append(
                    table_markdown([[cell.text for cell in row.cells] for row in block.rows])
                )
            else:
                style = block.style.name if block.style is not None else ""
                heading = re.fullmatch(r"Heading ([1-6])", style or "")
                prefix = "#" * int(heading[1]) + " " if heading else ""
                result.append(prefix + block.text)
        return "\n\n".join(result)

    def excel(self, path: Path) -> str:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        result = []
        try:
            for sheet in workbook.worksheets:
                rows = []
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= self.options.max_table_rows:
                        raise ValueError("Excel sheet exceeds row limit")
                    rows.append(["" if value is None else str(value) for value in row])
                result.append(f"## {sheet.title}\n\n{table_markdown(rows)}")
        finally:
            workbook.close()
        return "\n\n".join(result)

    def powerpoint(self, path: Path) -> str:
        from pptx import Presentation
        from pptx.shapes.graphfrm import GraphicFrame

        result = []
        for index, slide in enumerate(Presentation(str(path)).slides, start=1):
            content = [f"## Slide {index}"]
            for shape in slide.shapes:
                if isinstance(shape, GraphicFrame) and shape.has_table:
                    table = shape.table
                    if len(table.rows) > self.options.max_table_rows:
                        raise ValueError("PowerPoint table exceeds row limit")
                    content.append(
                        table_markdown([[cell.text for cell in row.cells] for row in table.rows])
                    )
                elif shape.has_text_frame:
                    content.append(str(getattr(shape, "text", "")))
            result.append("\n\n".join(content))
        return "\n\n".join(result)

    def image(self, path: Path) -> str:
        import pytesseract
        from PIL import Image, ImageOps

        pytesseract.pytesseract.tesseract_cmd = self.options.tesseract_command
        with Image.open(path) as image:
            if image.width * image.height > 50_000_000:
                raise ValueError("Image exceeds 50 megapixels")
            normalized = ImageOps.autocontrast(ImageOps.grayscale(ImageOps.exif_transpose(image)))
            return str(
                pytesseract.image_to_string(
                    normalized,
                    lang=self.options.ocr_language,
                    config="--psm 6",
                    timeout=self.options.ocr_timeout_seconds,
                )
            )
